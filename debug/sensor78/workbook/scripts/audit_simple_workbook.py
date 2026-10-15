"""Read-only independent audit of the compact detector-78 workbook.

Run after native Excel recalculation and save. This script never writes the XLSX.
It compares native cached outputs with the independent Python reference, checks
formula preservation, and audits package links and formula errors.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from workbook_paths import validate_output_paths


HERE = Path(__file__).resolve().parent
DEFAULT_BOOK = HERE.parent / "Sensor78_Emission_Comparison_Simple.xlsx"
DEFAULT_INPUTS = HERE.parents[1] / "results/simple_workbook_inputs.json"
DEFAULT_REPORT = HERE.parents[2] / "output/sensor78/workbook/current_workbook_audit.json"
SHEETS = [
    "01 Traffic inputs", "02 Emission rates", "03 Recorded O", "04 Mean speed S",
    "05 Reconstructed R", "06 Two-speed M", "07 Cubic A", "08 Results",
]
METHODS = ["O", "S", "R", "M", "A"]
POLLUTANTS = ["CO2", "NOx", "CO", "HC"]
METHOD_KEYS = {
    "O": "E_bin_g", "S": "E_avg_g", "R": "E_rec_g",
    "M": "E_2s_g", "A": "E_2s_cubic_g",
}
ERRORS = {
    "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!",
    "#SPILL!", "#CALC!", "#GETTING_DATA", "#CONNECT!", "#BLOCKED!", "#FIELD!",
}
EXTERNAL_FORMULA = re.compile(r"\[[^\]]+\](?:[^!]+!)")
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_BOOK)
    parser.add_argument("--inputs", type=Path, default=DEFAULT_INPUTS)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    validate_output_paths([("report", args.report)],
                          [("input workbook", args.workbook), ("input JSON", args.inputs)])
    path = args.workbook.resolve()
    source = json.loads(args.inputs.read_text())
    baseline = source["validation_only"]["expected_state"]
    failures = []
    if not path.is_file():
        raise SystemExit(f"Workbook is not ready: {path}")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cache = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if formulas.sheetnames != SHEETS:
        failures.append({"check": "sheet_order", "expected": SHEETS, "actual": formulas.sheetnames})

    counts = {}
    error_cells, blank_formula_caches, external_formulas = [], [], []
    for ws in formulas.worksheets:
        cached = cache[ws.title]
        total = 0
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                cv = cached[cell.coordinate].value
                if cell.data_type == "f":
                    total += 1
                    if cv is None:
                        blank_formula_caches.append(f"{ws.title}!{cell.coordinate}")
                    if isinstance(value, str) and EXTERNAL_FORMULA.search(value):
                        external_formulas.append({"cell": f"{ws.title}!{cell.coordinate}", "formula": value})
                if cached[cell.coordinate].data_type == "e" or (isinstance(cv, str) and cv in ERRORS):
                    error_cells.append({"cell": f"{ws.title}!{cell.coordinate}", "value": cv})
        counts[ws.title] = total
    if error_cells:
        failures.append({"check": "formula_errors", "count": len(error_cells)})
    if blank_formula_caches:
        failures.append({"check": "missing_formula_caches", "count": len(blank_formula_caches)})
    if external_formulas:
        failures.append({"check": "external_formula_references", "count": len(external_formulas)})

    totals, time_rows = [], []
    if "08 Results" in formulas.sheetnames:
        rs, rc = formulas["08 Results"], cache["08 Results"]
        for i, method in enumerate(METHODS, 7):
            for col, pollutant in enumerate(POLLUTANTS, 3):
                cell = rc.cell(i, col)
                target = baseline["rates"][pollutant][METHOD_KEYS[method]]
                actual = cell.value
                relative = abs(actual - target) / abs(target) if number(actual) else None
                live = rs.cell(i, col).data_type == "f"
                match = number(actual) and math.isclose(actual, target, rel_tol=1e-9, abs_tol=1e-7)
                totals.append({"method": method, "pollutant": pollutant,
                               "cell": f"08 Results!{cell.coordinate}", "expected_g": target,
                               "actual_g": actual, "relative_error": relative,
                               "formula_preserved": live, "matches": match})
                if not match or not live:
                    failures.append({"check": "emission_total", "method": method, "pollutant": pollutant,
                                     "matches": match, "formula_preserved": live})

    for sheet, method in zip(SHEETS[2:7], METHODS):
        if sheet not in formulas.sheetnames:
            continue
        ws, cached = formulas[sheet], cache[sheet]
        # The delivered workbook has a compact O/S layout. Find the labels
        # instead of assuming the generator's original table row numbers.
        header_rows = [r for r in range(1, ws.max_row + 1)
                       if ws.cell(r, 1).value == "Clock time"]
        total_rows = [r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 1).value == "Total / mean"]
        if len(header_rows) != 1 or len(total_rows) != 1:
            failures.append({"check": "time_series_labels", "sheet": sheet,
                             "header_rows": header_rows, "total_rows": total_rows})
            continue
        first_row, total_row = header_rows[0] + 1, total_rows[0]
        last_row = first_row + len(source["profile_rows"]) - 1
        active_rows = [r for r in range(first_row, last_row + 1)
                       if any(cached.cell(r, c).value is not None for c in range(1, 5))]
        row_formulas = [sum(ws.cell(r, c).data_type == "f" for c in range(1, ws.max_column + 1))
                        for r in range(first_row, last_row + 1)]
        total_formulas = sum(ws.cell(total_row, c).data_type == "f" for c in range(1, ws.max_column + 1))
        details = {"method": method, "sheet": sheet, "populated_time_rows": len(active_rows),
                   "expected_time_rows": len(source["profile_rows"]),
                   "first_data_row": first_row, "last_data_row": last_row,
                   "total_row": total_row, "minimum_formulas_per_time_row": min(row_formulas),
                   "formulas_in_total_row": total_formulas}
        time_rows.append(details)
        if len(active_rows) != len(source["profile_rows"]) or min(row_formulas) == 0 or total_formulas == 0:
            failures.append({"check": "time_series_table", **details})

    with zipfile.ZipFile(path) as package:
        files = package.namelist()
        external_parts = [name for name in files if name.startswith("xl/externalLinks/")]
        connection_parts = [name for name in files if name == "xl/connections.xml" or name.startswith("xl/queryTables/")]
        external_relationships = []
        for name in files:
            if not name.endswith(".rels"):
                continue
            tree = ET.fromstring(package.read(name))
            for relation in tree.findall(f"{{{REL_NS}}}Relationship"):
                if relation.attrib.get("TargetMode") == "External":
                    external_relationships.append({"part": name, **relation.attrib})
    # External hyperlinks, if any, are provenance links, not workbook dependencies.
    dependency_relationships = [r for r in external_relationships
                                if not r.get("Type", "").endswith("/hyperlink")]
    if external_parts or connection_parts or dependency_relationships:
        failures.append({"check": "standalone_package", "external_parts": len(external_parts),
                         "connection_parts": len(connection_parts),
                         "dependency_relationships": len(dependency_relationships)})
    formulas.close()
    cache.close()
    after = hashlib.sha256(path.read_bytes()).hexdigest()
    if before != after:
        failures.append({"check": "read_only_hash", "before": before, "after": after})
    report = {
        "workbook": path.name, "sha256": after, "read_only_hash_unchanged": before == after,
        "passed": not failures, "sheet_count": len(counts), "sheet_order_matches": list(counts) == SHEETS,
        "formula_count": sum(counts.values()), "formula_count_by_sheet": counts,
        "native_cached_values_present_for_all_formulas": not blank_formula_caches,
        "formula_error_count": len(error_cells), "formula_errors": error_cells,
        "missing_formula_caches": blank_formula_caches, "external_formulas": external_formulas,
        "external_link_parts": external_parts, "connection_parts": connection_parts,
        "external_dependency_relationships": dependency_relationships,
        "external_hyperlinks": [r for r in external_relationships if r not in dependency_relationships],
        "time_series_checks": time_rows, "emission_total_checks": totals,
        "all_20_emission_totals_match": len(totals) == 20 and all(r["matches"] for r in totals),
        "max_total_relative_error": max((r["relative_error"] for r in totals if r["relative_error"] is not None), default=None),
        "failures": failures,
        "scope": "Independent read-only cached-output and formula audit. Native recalculation and input-perturbation evidence must be supplied separately; this audit does not infer those actions from cached values.",
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2, default=str, allow_nan=False))
    print(json.dumps({k: report[k] for k in ["passed", "sheet_count", "formula_count",
                     "formula_error_count", "all_20_emission_totals_match", "max_total_relative_error", "failures"]}, indent=2))
    print(f"Audit report: {args.report.resolve()}")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
